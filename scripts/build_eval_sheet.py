"""Pool retrieval candidates for the eval queries and emit an expert-review workbook.

Produces a multi-tab .xlsx (one tab per query) that a hadith scholar can fill
in with 0-3 relevance judgments. Upload it to Google Drive and it converts to
a Google Sheet with the tabs intact.

Pipeline per query:
  1. retrieve_union_multi([query] + hand-authored variants) -> RRF-merged pool
  2. cross-encoder rerank the pool
  3. force-include any `seeds` (refs already vetted during interactive curation)
  4. truncate to --per-query rows total, seeds first

The `variants` are static, hand-authored, and live in the queries JSON. No LLM
is called here: the 2026-05-29 session showed that letting a model supply
hadith NUMBERS produces confident hallucinations. Variants only widen the
*query phrasing*, so a bad variant costs recall, never correctness.

Usage:
    python -m scripts.build_eval_sheet
    python -m scripts.build_eval_sheet --per-query 15 --pool-k 30
    python -m scripts.build_eval_sheet --no-rerank      # skip the cross-encoder
"""

from __future__ import annotations

import argparse
import json
import time
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from sunnah_toolkit.core import reranker as reranker_mod
from sunnah_toolkit.core.data import Hadith, Library, load, strip_narrator_markup
from sunnah_toolkit.core.retrieval import Candidate, retrieve_union_multi

ROOT = Path(__file__).resolve().parent.parent
QUERIES_PATH = ROOT / "docs" / "eval" / "queries-20.json"
OUT_DIR = ROOT / "docs" / "eval"

TEXT_TRUNCATE = 600
BLANK_ADD_ROWS = 8

HEADERS = [
    "#", "Relevance (0-3)", "Notes",
    "Collection", "Hadith #", "Link", "Grade", "Narrator",
    "English", "Arabic", "Found by", "Rank",
]
COL_WIDTHS = [4, 15, 28, 15, 10, 34, 14, 22, 70, 60, 20, 6]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SEED_FILL = PatternFill("solid", fgColor="E2EFDA")
BANNER_FONT = Font(bold=True, size=13)
ADD_FILL = PatternFill("solid", fgColor="FCE4D6")

SCALE_TEXT = [
    ("3", "Primary / canonical", "A scholar answering this query would cite this hadith."),
    ("2", "Directly relevant", "Clearly on-topic; belongs in the answer set."),
    ("1", "Marginal", "Touches the topic but is not what was asked."),
    ("0", "Not relevant", "Off-topic; the search engine should not have returned it."),
]


def _truncate(text: str, n: int = TEXT_TRUNCATE) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _cite(h: Hadith) -> str:
    """sunnah.com citation id: first of a paired range, whitespace squashed."""
    raw = h.hadith_number or str(h.id_in_book)
    return "".join(raw.split(",", 1)[0].split())


def _link(h: Hadith) -> str:
    return f"https://sunnah.com/{h.collection}:{_cite(h)}"


def _score_pool(
    library: Library, query: str, candidates: list[Candidate], use_rerank: bool
) -> list[tuple[Candidate, float]]:
    """Rerank the pool, or fall back to RRF order with descending pseudo-scores."""
    if not candidates:
        return []
    if not use_rerank or not reranker_mod.reranker_enabled():
        return [(c, float(len(candidates) - i)) for i, c in enumerate(candidates)]

    name = reranker_mod.default_reranker_name()
    r = reranker_mod.get_reranker(name)
    docs = []
    for c in candidates:
        h = c.hadith
        col = library.get_collection(h.collection)
        title = col.english_title if col else h.collection
        docs.append(
            f"{title}\n{h.english_narrator}\n{h.english_text}\n"
            f"{strip_narrator_markup(h.arabic)}"
        )
    scores = r.score(query, docs)
    scored = list(zip(candidates, scores))
    scored.sort(key=lambda p: p[1], reverse=True)
    return scored


def _resolve_seeds(library: Library, seeds: list[dict]) -> list[tuple[Hadith, str]]:
    """Resolve vetted refs through the fixed hadith_number lookup.

    Anything that fails to resolve is reported loudly rather than silently
    dropped -- a dead seed means the curation notes carry a bad reference.
    """
    out: list[tuple[Hadith, str]] = []
    for s in seeds:
        col, num = s["collection"], s["hadith_number"]
        h = library.get_hadith(col, num)
        if h is None:
            print(f"    !! seed {col}:{num} did not resolve -- SKIPPED")
            continue
        out.append((h, f"{col}:{num}"))
    return out


def build_pool(
    library: Library,
    q: dict[str, Any],
    per_query: int,
    pool_k: int,
    use_rerank: bool,
) -> list[dict[str, Any]]:
    queries = [q["query"], *q.get("variants", [])]
    t0 = time.perf_counter()
    candidates = retrieve_union_multi(queries, k_per_retriever=pool_k)
    scored = _score_pool(library, q["query"], candidates, use_rerank)

    seeds = _resolve_seeds(library, q.get("seeds", []))
    seed_urns = {h.urn_arabic for h, _ in seeds}

    rows: list[dict[str, Any]] = []
    seen: set[int] = set()

    # Seeds first -- already vetted, so they never lose a slot to a system hit.
    sys_rank = {c.hadith.urn_arabic: i + 1 for i, (c, _) in enumerate(scored)}
    sys_srcs = {c.hadith.urn_arabic: sorted(c.sources) for c, _ in scored}
    for h, _ref in seeds:
        seen.add(h.urn_arabic)
        found = ["seed"] + sys_srcs.get(h.urn_arabic, [])
        rows.append({
            "hadith": h,
            "found_by": ", ".join(found),
            "rank": sys_rank.get(h.urn_arabic, ""),
            "is_seed": True,
        })

    for c, _score in scored:
        if len(rows) >= per_query:
            break
        h = c.hadith
        if h.urn_arabic in seen or h.urn_arabic in seed_urns:
            continue
        seen.add(h.urn_arabic)
        rows.append({
            "hadith": h,
            "found_by": ", ".join(sorted(c.sources)),
            "rank": sys_rank.get(h.urn_arabic, ""),
            "is_seed": False,
        })

    rows = rows[:per_query]
    elapsed = time.perf_counter() - t0
    n_seed = sum(1 for r in rows if r["is_seed"])
    print(
        f"  {q['id']} {q['query'][:44]:<46} pool={len(candidates):>4} "
        f"rows={len(rows):>2} (seeds={n_seed}) {elapsed:>5.1f}s"
    )
    return rows


def _write_readme(wb: Workbook, meta: dict, n_queries: int, per_query: int) -> None:
    ws = wb.create_sheet("README", 0)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 88

    def row(a="", b="", c="", bold=False, size=11):
        ws.append([a, b, c])
        if bold:
            for col in "ABC":
                ws[f"{col}{ws.max_row}"].font = Font(bold=True, size=size)
        ws[f"C{ws.max_row}"].alignment = Alignment(wrap_text=True, vertical="top")

    row("Hadith search — expert relevance review", bold=True, size=14)
    row()
    row("What this is", "", "Each tab is one search query. The rows are hadiths our search engine "
        "returned for that query. We need a scholar to tell us which ones are actually "
        "the right answer, so we can measure and improve the ranking.")
    row()
    row("What to do", "", "In every tab, fill the yellow 'Relevance (0-3)' column for each row. "
        "Use the dropdown. 'Notes' is optional and free-form — corrections, caveats, "
        "and objections are all welcome there.")
    row()
    row("The scale", bold=True)
    for score, label, desc in SCALE_TEXT:
        row(score, label, desc)
    row()
    row("Missing hadiths", "", "The engine is imperfect and will have missed things. At the bottom of "
        "each tab there is an orange 'ADD MISSING HADITHS BELOW' block. Paste a "
        "sunnah.com link (e.g. https://sunnah.com/bukhari:6114) into the Link column "
        "and set its relevance. This is some of the most valuable feedback you can give — "
        "it tells us what the engine cannot currently find at all.")
    row()
    row("Green rows", "", "Rows shaded green were pre-selected by a non-specialist during earlier "
        "drafting. Please judge them exactly as critically as the rest — several "
        "earlier picks turned out to be wrong.")
    row()
    row("Splitting the work", "", "The Index tab has a Reviewer column. Multiple people can take "
        "different queries; note your name there so work isn't duplicated.")
    row()
    row("Please don't", "", "Reorder rows, rename tabs, or delete columns — the results are read back "
        "programmatically and matched on Collection + Hadith #.")
    row()
    row("Generated", "", f"{meta['generated']} · {n_queries} queries · up to {per_query} candidates each "
        f"· reranker: {meta['reranker']}")


def _write_index(wb: Workbook, entries: list[dict]) -> None:
    ws = wb.create_sheet("Index", 1)
    headers = ["Query ID", "Query", "Mode", "Tab name", "Candidates", "Reviewer", "Status"]
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for w, col in zip([10, 46, 12, 26, 12, 20, 16], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    for e in entries:
        ws.append([e["id"], e["query"], e["mode_hint"], e["tab"], e["n"], "", "Not started"])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=6).fill = INPUT_FILL
        ws.cell(row=r, column=7).fill = INPUT_FILL
    ws.freeze_panes = "A2"


def _write_query_tab(wb: Workbook, q: dict, rows: list[dict], tab: str) -> None:
    ws = wb.create_sheet(tab)

    ws.append([f"{q['id']}  ·  {q['query']}"])
    ws["A1"].font = BANNER_FONT
    ws.append([f"mode: {q['mode_hint']}   ·   variants searched: "
               + " | ".join(q.get("variants", []))])
    ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws.append([])

    hdr_row = 4
    ws.append(HEADERS)
    for i, _ in enumerate(HEADERS, start=1):
        c = ws.cell(row=hdr_row, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for n, r in enumerate(rows, start=1):
        h: Hadith = r["hadith"]
        ws.append([
            n, "", "",
            h.collection, _cite(h), _link(h),
            h.english_grade or "—",
            _truncate(h.english_narrator, 120),
            _truncate(h.english_text),
            _truncate(strip_narrator_markup(h.arabic)),
            r["found_by"], r["rank"],
        ])
        excel_row = ws.max_row
        ws.cell(row=excel_row, column=2).fill = INPUT_FILL
        ws.cell(row=excel_row, column=3).fill = INPUT_FILL
        if r["is_seed"]:
            for col in (4, 5, 6, 7, 8):
                ws.cell(row=excel_row, column=col).fill = SEED_FILL
        for col in (8, 9, 10):
            ws.cell(row=excel_row, column=col).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
        ws.cell(row=excel_row, column=10).alignment = Alignment(
            wrap_text=True, vertical="top", horizontal="right"
        )

    ws.append([])
    ws.append(["ADD MISSING HADITHS BELOW — paste a sunnah.com link in the Link column"])
    marker = ws.max_row
    ws.cell(row=marker, column=1).font = Font(bold=True)
    for col in range(1, len(HEADERS) + 1):
        ws.cell(row=marker, column=col).fill = ADD_FILL
    for _ in range(BLANK_ADD_ROWS):
        ws.append([""] * len(HEADERS))
        for col in (2, 3, 6):
            ws.cell(row=ws.max_row, column=col).fill = INPUT_FILL

    dv = DataValidation(
        type="list", formula1='"3,2,1,0"', allow_blank=True, showDropDown=False
    )
    dv.error = "Enter 3, 2, 1 or 0 (see the README tab)."
    dv.errorTitle = "Relevance must be 0-3"
    ws.add_data_validation(dv)
    dv.add(f"B{hdr_row + 1}:B{ws.max_row}")

    ws.freeze_panes = f"D{hdr_row + 1}"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--per-query", type=int, default=15,
                    help="Max candidate rows per query tab (default 15).")
    ap.add_argument("--pool-k", type=int, default=30,
                    help="Top-K per retriever per variant (default 30).")
    ap.add_argument("--no-rerank", action="store_true",
                    help="Skip the cross-encoder; keep RRF order.")
    args = ap.parse_args()

    use_rerank = not args.no_rerank
    spec = json.loads(QUERIES_PATH.read_text())
    queries = spec["queries"]

    print(f"Loading library ...")
    library = load()
    reranker_name = (
        reranker_mod.default_reranker_name() if use_rerank else "none (RRF order)"
    )
    print(f"Pooling {len(queries)} queries · per_query={args.per_query} · "
          f"pool_k={args.pool_k} · reranker={reranker_name}\n")

    entries: list[dict] = []
    pools: dict[str, list[dict]] = {}
    t0 = time.perf_counter()
    for q in queries:
        rows = build_pool(library, q, args.per_query, args.pool_k, use_rerank)
        pools[q["id"]] = rows
        entries.append({
            "id": q["id"], "query": q["query"], "mode_hint": q["mode_hint"],
            "tab": f"{q['id']}_{q['slug']}"[:31], "n": len(rows),
        })
    print(f"\nPooling done in {time.perf_counter() - t0:.1f}s")

    stamp = time.strftime("%Y%m%d")
    meta = {"generated": time.strftime("%Y-%m-%d %H:%M"), "reranker": reranker_name}

    wb = Workbook()
    wb.remove(wb.active)
    _write_readme(wb, meta, len(queries), args.per_query)
    _write_index(wb, entries)
    for q in queries:
        _write_query_tab(wb, q, pools[q["id"]], f"{q['id']}_{q['slug']}"[:31])

    xlsx_path = OUT_DIR / f"expert-sheet-{stamp}.xlsx"
    wb.save(xlsx_path)
    print(f"Wrote {xlsx_path}  ({xlsx_path.stat().st_size // 1024} KB)")

    pool_path = OUT_DIR / f"pool-{stamp}.json"
    pool_path.write_text(json.dumps({
        "_generated": meta["generated"],
        "_reranker": reranker_name,
        "_per_query": args.per_query,
        "_pool_k": args.pool_k,
        "queries": [
            {
                "id": q["id"], "query": q["query"], "mode_hint": q["mode_hint"],
                "tab": f"{q['id']}_{q['slug']}"[:31],
                "candidates": [
                    {
                        "n": i + 1,
                        "collection": r["hadith"].collection,
                        "hadith_number": _cite(r["hadith"]),
                        "urn": r["hadith"].urn_arabic,
                        "grade": r["hadith"].english_grade,
                        "found_by": r["found_by"],
                        "rank": r["rank"],
                        "is_seed": r["is_seed"],
                    }
                    for i, r in enumerate(pools[q["id"]])
                ],
            }
            for q in queries
        ],
    }, indent=2, ensure_ascii=False))
    print(f"Wrote {pool_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
